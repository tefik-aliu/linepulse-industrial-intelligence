from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import build_overview
from .models import ProductionEvent


def make_shift_workbook(events: list[ProductionEvent], dependency_degraded: bool) -> bytes:
    overview = build_overview(events, dependency_degraded)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Shift Summary"

    title_fill = PatternFill("solid", fgColor="0B1E34")
    accent_fill = PatternFill("solid", fgColor="1B6CF2")
    light_fill = PatternFill("solid", fgColor="DCEBFF")

    summary["A1"] = "LinePulse Shift Intelligence Report"
    summary["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    summary["A1"].fill = title_fill
    summary.merge_cells("A1:D1")

    summary.append(["Metric", "Value", "Operational meaning", "Status"])
    for cell in summary[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = accent_fill

    kpis = overview["kpis"]
    rows = [
        (
            "OEE",
            f"{kpis['oee']}%",
            "Combined availability, performance and quality",
            "Review" if kpis["oee"] < 70 else "Healthy",
        ),
        (
            "Availability",
            f"{kpis['availability']}%",
            "Time producing versus planned time",
            "Review" if kpis["availability"] < 85 else "Healthy",
        ),
        (
            "Performance",
            f"{kpis['performance']}%",
            "Actual throughput versus ideal cycle",
            "Review" if kpis["performance"] < 90 else "Healthy",
        ),
        (
            "Quality",
            f"{kpis['quality']}%",
            "Good output versus total output",
            "Review" if kpis["quality"] < 97 else "Healthy",
        ),
        (
            "Scrap rate",
            f"{kpis['scrap_rate']}%",
            "Rejected units versus total output",
            "Review" if kpis["scrap_rate"] > 3 else "Healthy",
        ),
        (
            "Downtime",
            f"{kpis['downtime_minutes']} min",
            "Accumulated recorded downtime",
            "Observe",
        ),
        (
            "MTBF",
            f"{kpis['mtbf_minutes']} min",
            "Mean run time between material failures",
            "Observe",
        ),
        (
            "MTTR",
            f"{kpis['mttr_minutes']} min",
            "Mean time to repair recorded failures",
            "Observe",
        ),
    ]
    for row in rows:
        summary.append(row)

    alert_start = summary.max_row + 2
    summary.cell(alert_start, 1, "Active explainable alerts")
    summary.cell(alert_start, 1).font = Font(bold=True, color="FFFFFF")
    summary.cell(alert_start, 1).fill = title_fill
    summary.merge_cells(start_row=alert_start, start_column=1, end_row=alert_start, end_column=4)

    summary.append(["Severity", "Title", "Evidence", "Recommended action"])
    for cell in summary[alert_start + 1]:
        cell.font = Font(bold=True)
        cell.fill = light_fill

    alerts = overview["alerts"] or [
        {
            "severity": "info",
            "title": "No active threshold alerts",
            "evidence": "Current rolling windows remain within configured limits.",
            "recommendation": "Continue monitoring.",
        }
    ]
    for alert in alerts:
        summary.append(
            [
                alert["severity"].upper(),
                alert["title"],
                alert["evidence"],
                alert["recommendation"],
            ]
        )

    events_sheet = workbook.create_sheet("Production Events")
    event_headers = [
        "Timestamp", "Machine", "Shift", "Product", "Good", "Rejects", "Downtime s",
        "Stop reason", "Defect", "Pressure bar", "Cycle s", "Scenario"
    ]
    events_sheet.append(event_headers)
    for cell in events_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = accent_fill

    for event in reversed(events[-100:]):
        events_sheet.append(
            [
                event.timestamp.isoformat(),
                event.machine,
                event.shift,
                event.product,
                event.good_units,
                event.reject_units,
                round(event.downtime_seconds, 1),
                event.stop_reason,
                event.defect_type,
                round(event.cutting_pressure_bar, 2),
                round(event.cycle_time_seconds, 2),
                event.scenario,
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A3" if sheet.title == "Shift Summary" else "A2"
        for column_index in range(1, sheet.max_column + 1):
            values = [
                sheet.cell(row=row_index, column=column_index).value
                for row_index in range(1, sheet.max_row + 1)
            ]
            max_length = max(len(str(value or "")) for value in values)
            width = min(max(max_length + 2, 12), 46)
            sheet.column_dimensions[get_column_letter(column_index)].width = width
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
